import io
import re

def slugify(text: str) -> str:
    text = re.sub(r"[^\w\s-]", "", text or "").strip().lower()
    return re.sub(r"[\s_-]+", "-", text) or "document"

def _lines(body):
    if body is None:
        return []
    if isinstance(body, list):
        body = "\n".join(str(item) for item in body)
    else:
        body = str(body)
    return [ln.strip() for ln in body.split("\n") if ln.strip()]

def build_docx(title: str, sections: list) -> io.BytesIO:
    from docx import Document

    doc = Document()
    doc.add_heading(str(title or "Document"), level=0)

    for sec in sections:
        heading = str(sec.get("heading") or "").strip()
        if heading:
            doc.add_heading(heading, level=1)
        for line in _lines(sec.get("body")):
            if line.startswith("- "):
                doc.add_paragraph(line[2:], style="List Bullet")
            else:
                doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf

def build_pptx(title: str, sections: list) -> io.BytesIO:
    from pptx import Presentation

    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = str(title or "Presentation")
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = ""

    bullet_layout = prs.slide_layouts[1]
    for sec in sections:
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = str(sec.get("heading") or "").strip() or " "
        body_tf = slide.placeholders[1].text_frame
        body_tf.clear()
        lines = _lines(sec.get("body"))
        if not lines:
            continue
        for i, line in enumerate(lines):
            text = line[2:] if line.startswith("- ") else line
            if i == 0:
                body_tf.text = text
            else:
                p = body_tf.add_paragraph()
                p.text = text

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf

def _xml_escape(text) -> str:
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )

def build_pdf(title: str, sections: list) -> io.BytesIO:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem
    from reportlab.lib.styles import getSampleStyleSheet

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, topMargin=54, bottomMargin=54)
    styles = getSampleStyleSheet()
    story = [Paragraph(_xml_escape(title or "Document"), styles["Title"]), Spacer(1, 14)]

    for sec in sections:
        heading = str(sec.get("heading") or "").strip()
        if heading:
            story.append(Paragraph(_xml_escape(heading), styles["Heading2"]))

        bullets = []
        for line in _lines(sec.get("body")):
            if line.startswith("- "):
                bullets.append(ListItem(Paragraph(_xml_escape(line[2:]), styles["Normal"])))
            else:
                if bullets:
                    story.append(ListFlowable(bullets, bulletType="bullet", leftIndent=18))
                    bullets = []
                story.append(Paragraph(_xml_escape(line), styles["Normal"]))
        if bullets:
            story.append(ListFlowable(bullets, bulletType="bullet", leftIndent=18))
        story.append(Spacer(1, 12))

    doc.build(story)
    buf.seek(0)
    return buf

BUILDERS = {"docx": build_docx, "pptx": build_pptx, "pdf": build_pdf}
MIME_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
}
